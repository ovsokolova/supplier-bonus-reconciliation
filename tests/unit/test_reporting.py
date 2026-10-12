"""Тесты экспорта в Excel."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from supplier_bonus.reporting.excel import (
    SHEET_ISSUES,
    SHEET_ROWS,
    SHEET_SUMMARY,
    export_to_excel,
)
from supplier_bonus.schemas import (
    MatchStatus,
    ReconciliationResult,
    ReconciliationRow,
    ReconciliationSummary,
)


@pytest.fixture()
def sample_result() -> ReconciliationResult:
    return ReconciliationResult(
        summary=ReconciliationSummary(
            matched_by_inn=1,
            matched_within_tolerance=1,
            amount_mismatch=1,
            missing_in_actuals=1,
        ),
        rows=[
            ReconciliationRow(
                supplier_name="ООО «Ромашка»",
                period="2026-Q1",
                expected=100.0,
                actual=100.0,
                diff=0.0,
                status=MatchStatus.MATCHED_WITHIN_TOLERANCE,
            ),
            ReconciliationRow(
                supplier_name="ООО «Василёк»",
                period="2026-Q1",
                expected=250.0,
                actual=200.0,
                diff=50.0,
                status=MatchStatus.AMOUNT_MISMATCH,
            ),
            ReconciliationRow(
                supplier_name="ООО «Подсолнух»",
                period="2026-Q1",
                expected=75.0,
                actual=None,
                status=MatchStatus.MISSING_IN_ACTUALS,
            ),
        ],
        amount_tolerance=1.0,
        fuzzy_threshold=90,
    )


def test_export_creates_file(tmp_path: Path, sample_result: ReconciliationResult) -> None:
    out = tmp_path / "report.xlsx"
    export_to_excel(sample_result, out)
    assert out.exists()
    assert out.stat().st_size > 0


def test_export_has_three_sheets(tmp_path: Path, sample_result: ReconciliationResult) -> None:
    out = tmp_path / "report.xlsx"
    export_to_excel(sample_result, out)
    wb = load_workbook(out)
    assert SHEET_ROWS in wb.sheetnames
    assert SHEET_SUMMARY in wb.sheetnames
    assert SHEET_ISSUES in wb.sheetnames


def test_rows_sheet_contains_all_rows(tmp_path: Path, sample_result: ReconciliationResult) -> None:
    out = tmp_path / "report.xlsx"
    export_to_excel(sample_result, out)
    wb = load_workbook(out)
    ws = wb[SHEET_ROWS]
    # header + 3 rows
    assert ws.max_row == 4


def test_issues_sheet_filters_statuses(tmp_path: Path, sample_result: ReconciliationResult) -> None:
    out = tmp_path / "report.xlsx"
    export_to_excel(sample_result, out)
    wb = load_workbook(out)
    ws = wb[SHEET_ISSUES]
    # header + amount_mismatch + missing_in_actuals
    assert ws.max_row == 3


def test_summary_sheet_has_metrics(tmp_path: Path, sample_result: ReconciliationResult) -> None:
    out = tmp_path / "report.xlsx"
    export_to_excel(sample_result, out)
    wb = load_workbook(out)
    ws = wb[SHEET_SUMMARY]
    values = [ws.cell(row=i, column=1).value for i in range(1, ws.max_row + 1)]
    assert "amount_tolerance" in values
    assert "fuzzy_threshold" in values
